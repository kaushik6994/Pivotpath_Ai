"""Meaningful rule, file handling, and retrieval regressions; no network calls."""

from io import BytesIO
import unittest

from backend.document_io import export_pdf, extract_text
from backend.engine import analyze_document
from backend.knowledge import answer_question


def rule(**overrides):
    return {
        "id": "r1", "name": "Document owner", "type": "required_text",
        "value": "Owner:", "weight": 10, "category": "All",
        "enabled": True, "severity": "major", "suggestion": "Name the owner.",
    } | overrides


class RuleEngineTests(unittest.TestCase):
    def test_weighted_score_and_critical_gate(self):
        report = analyze_document("Owner: Quality", [rule(weight=90), rule(id="r2", value="Effective date", weight=10, severity="critical")], {"threshold": 85})
        self.assertEqual(report["score"], 90)
        self.assertEqual((report["passed"], report["failed"]), (1, 1))
        self.assertFalse(report["eligible"])

    def test_threshold_uses_unrounded_score(self):
        report = analyze_document("Owner:", [rule(weight=90), rule(id="r2", value="missing", weight=10.01)], {"threshold": 90})
        self.assertEqual(report["score"], 90)
        self.assertFalse(report["eligible"])

    def test_disabled_and_other_category_do_not_affect_score(self):
        report = analyze_document("OWNER:\nJane", [rule(), rule(id="disabled", value="absent", enabled=False), rule(id="other", value="absent", category="SOP")], {"threshold": 100}, "Quality")
        self.assertTrue(report["eligible"])
        self.assertEqual(report["total"], 1)

    def test_required_text_normalizes_unicode_case_and_whitespace(self):
        report = analyze_document("QUALITY\n  ASSURANCE", [rule(value="quality assurance")], {})
        self.assertTrue(report["findings"][0]["passed"])

    def test_required_sections_require_headings_not_prose_mentions(self):
        check = rule(type="required_sections", value="Purpose, Scope, Responsibilities")
        report = analyze_document("# Purpose\n1. Scope\nWe discuss responsibilities here.", [check], {})
        self.assertFalse(report["findings"][0]["passed"])
        self.assertIn("Responsibilities", report["findings"][0]["evidence"])
        report = analyze_document("# Purpose\n1. Scope\n## Responsibilities:", [check], {})
        self.assertTrue(report["eligible"])

    def test_prohibited_text_is_literal_not_regex(self):
        check = rule(type="prohibited_text", value=".*")
        self.assertTrue(analyze_document("Safe wording.", [check], {})["eligible"])
        self.assertFalse(analyze_document("Expression .* is disallowed.", [check], {})["eligible"])

    def test_word_count_supports_unicode_and_hyphenated_words(self):
        report = analyze_document("The café uses well-tested equipment.", [rule(type="min_words", value="5")], {})
        self.assertTrue(report["eligible"])
        self.assertIn("5 words", report["findings"][0]["evidence"])

    def test_empty_rulebook_never_eligible(self):
        report = analyze_document("Content", [], {"threshold": 0})
        self.assertEqual(report["score"], 0)
        self.assertFalse(report["eligible"])

    def test_invalid_rule_cannot_accidentally_pass(self):
        for invalid in [rule(type="min_words", value="-1"), rule(weight=float("nan")), rule(type="unknown")]:
            report = analyze_document("Owner:", [invalid], {})
            self.assertFalse(report["eligible"])
            self.assertEqual(report["score"], 0)

    def test_learning_only_affects_suggestions_with_approval(self):
        feedback = [{"rule_id": "r1", "status": "approved", "correction": "Add the accountable owner.", "created_at": "2026-10-01"}, {"rule_id": "r1", "status": "pending", "correction": "Bypass rule.", "created_at": "2026-10-02"}]
        report = analyze_document("Missing attribution", [rule()], {"learning_enabled": True}, approved_feedback=feedback)
        self.assertEqual(report["findings"][0]["suggestion"], "Add the accountable owner.")
        self.assertFalse(report["eligible"])
        disabled = analyze_document("Missing attribution", [rule()], {"learning_enabled": False}, approved_feedback=feedback)
        self.assertEqual(disabled["findings"][0]["suggestion"], "Name the owner.")


class DocumentIOTests(unittest.TestCase):
    def test_pdf_export_round_trip_escapes_markup_and_unicode(self):
        data = export_pdf("Quality & Safety <review>", "Purpose\nCafé · β = 2\n<script>alert('x')</script>", {"version": 2})
        self.assertTrue(data.startswith(b"%PDF-"))
        text = extract_text("report.pdf", data)
        self.assertIn("Café", text)
        self.assertIn("β", text)
        self.assertIn("<script>", text)
        self.assertIn("Version 2", text)

    def test_text_encoding_and_binary_rejection(self):
        self.assertEqual(extract_text("notes.md", b"\xef\xbb\xbfHello\r\nWorld"), "Hello\nWorld")
        self.assertEqual(extract_text("notes.txt", "Café".encode("utf-16")), "Café")
        with self.assertRaisesRegex(ValueError, "binary"):
            extract_text("notes.txt", b"hello\x00binary")
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            extract_text("notes.txt", b"\xffbad")

    def test_invalid_and_oversize_inputs(self):
        for filename, data, message in [("file.exe", b"abc", "Supported"), ("file.pdf", b"", "empty"), ("file.pdf", b"not pdf", "header"), ("file.txt", b"x" * (10 * 1024 * 1024 + 1), "10 MB")]:
            with self.subTest(filename=filename, message=message), self.assertRaisesRegex(ValueError, message):
                extract_text(filename, data)

    def test_scanned_and_encrypted_pdfs_rejected(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=600, height=800)
        blank = BytesIO()
        writer.write(blank)
        with self.assertRaisesRegex(ValueError, "OCR"):
            extract_text("scan.pdf", blank.getvalue())
        writer.encrypt("secret")
        encrypted = BytesIO()
        writer.write(encrypted)
        with self.assertRaisesRegex(ValueError, "password"):
            extract_text("encrypted.pdf", encrypted.getvalue())

    def test_docx_text_and_table_order(self):
        from docx import Document

        document = Document()
        document.add_paragraph("Before table")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Batch"
        table.cell(0, 1).text = "A-42"
        document.add_paragraph("After table")
        output = BytesIO()
        document.save(output)
        text = extract_text("batch.docx", output.getvalue())
        self.assertLess(text.index("Before table"), text.index("Batch"))
        self.assertLess(text.index("Batch"), text.index("After table"))


class KnowledgeTests(unittest.TestCase):
    documents = [
        {"id": "cold", "title": "Cold chain SOP", "version": 3, "content": "Temperature monitoring\nStore the product between 2 and 8 degrees Celsius. Record storage temperature every four hours."},
        {"id": "training", "title": "Training policy", "version": 2, "content": "All employees complete annual safety training. Attendance records are retained by Human Resources."},
    ]

    def test_retrieves_correct_document_and_version_with_quotes(self):
        answer = answer_question("What is the storage temperature?", self.documents)
        self.assertEqual(answer["mode"], "extractive")
        self.assertEqual(answer["sources"][0]["document_id"], "cold")
        self.assertEqual(answer["sources"][0]["version"], 3)
        self.assertIn("2 and 8", answer["answer"])
        self.assertIn("[1]", answer["answer"])

    def test_absent_evidence_does_not_invent_answer(self):
        answer = answer_question("What is the approved spaceship propulsion?", self.documents)
        self.assertEqual(answer["sources"], [])
        self.assertIn("could not find", answer["answer"])

    def test_filtered_documents_are_the_only_possible_sources(self):
        answer = answer_question("What is the storage temperature?", self.documents[1:])
        self.assertEqual(answer["sources"], [])

    def test_empty_and_long_questions(self):
        with self.assertRaisesRegex(ValueError, "Enter a question"):
            answer_question("  ", self.documents)
        with self.assertRaisesRegex(ValueError, "2,000"):
            answer_question("x" * 2001, self.documents)
        self.assertEqual(answer_question("please tell me", self.documents)["sources"], [])

    def test_form_feed_preserves_extracted_source_page(self):
        answer = answer_question("temperature", [{"id": "p", "title": "Pages", "version": 1, "content": "Administrative information.\fStorage temperature is monitored."}])
        self.assertEqual(answer["sources"][0]["page"], 2)


if __name__ == "__main__":
    unittest.main()
